#!/usr/bin/env python3
"""Regression checks for review.py, run on a throwaway copy of the pours folder (the real one is never touched).

Run with the studio's venv python (it has openpyxl):
  Dionysus/Cocktail_Website_Agent/design-artifacts/pours/_studio/.venv/bin/python skills/dps-review-desk/scripts/tests/test_review.py
"""
import hashlib, importlib, json, os, re, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "..", "review.py")
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "dps-tools", "scripts"))
import _common  # noqa: E402

REAL_POURS = _common.POURS
TMP = tempfile.mkdtemp(prefix="dps-review-test-")
POURS = os.path.join(TMP, "pours")
shutil.copytree(REAL_POURS, POURS, ignore=shutil.ignore_patterns(".venv", "work", "review*.xlsx", "~$*", "review-archive",
                                                                  "desk-notes.json", "robin-edits.md"))
os.environ["DPS_POURS_FOLDER"] = POURS
importlib.reload(_common)   # every module imported from here on points at the copy
sys.path.insert(0, os.path.join(HERE, ".."))
import review  # noqa: E402
import pourfile  # noqa: E402
import openpyxl  # noqa: E402

STUDIO = os.path.join(POURS, "_studio")
if os.path.realpath(review.POURS) != os.path.realpath(POURS):
    sys.exit("refusing to run: review.py points at %s, not the test copy" % review.POURS)
fails = 0


def check(name, ok, detail=""):
    global fails
    print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  (%s)" % detail) if detail and not ok else ""))
    fails += not ok


def run(*args):
    p = subprocess.run([sys.executable, SCRIPT] + list(args), capture_output=True, text=True, env=dict(os.environ))
    try:
        return p.returncode, json.loads(p.stdout)
    except ValueError:
        return p.returncode, {"raw": p.stdout + p.stderr}


def hashes():
    return {f: hashlib.sha1(open(f, "rb").read()).hexdigest() for f in pourfile.all_pours(POURS)}


def edit(values):
    """Play Robin: {pairing: {column: value}} into the Pours sheet (and 'rules' for the rule sheet)."""
    wb = openpyxl.load_workbook(review.WORKBOOK)
    ws = wb["Pours"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        p = row[head.index("pairing")].value
        for col, v in values.get(p, {}).items():
            row[head.index(col)].value = v
    if "rules" in values:
        rs = wb["Rule candidates"]
        rh = [c.value for c in rs[1]]
        for row in rs.iter_rows(min_row=2):
            for col, v in values["rules"].items():
                row[rh.index(col)].value = v
    wb.save(review.WORKBOOK)


def cell(pairing, col):
    wb = openpyxl.load_workbook(review.WORKBOOK)
    ws = wb["Pours"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[head.index("pairing")] == pairing:
            return row[head.index(col)]


def set_status(pairing, status):
    p = os.path.join(POURS, pairing + ".md")
    t = open(p, encoding="utf-8").read()
    open(p, "w", encoding="utf-8").write(re.sub(r"^status:.*$", "status: " + status, t, count=1, flags=re.M))


try:
    # 1. the writer is exact: writing every field back unchanged leaves each pour byte-identical
    ok = True
    for f in pourfile.all_pours(POURS):
        d = pourfile.parse(f)
        n = len(d["yours"]) + 1
        cells = review.cells_from(d, n)
        ch = dict(cells, _yours=review.yours_from(cells, n))
        text = open(f, encoding="utf-8").read()
        if review.apply_fields(text, ch, n) != text:
            ok = False
            print("   writer changed", os.path.basename(f))
    check("writer: every field written back unchanged = identical file (all pours)", ok)

    # 2. export
    for p in ("sage-lover", "magician-outlaw", "creator-hero"):
        set_status(p, "draft")
    before = hashes()
    rc, out = run("export")
    wb = openpyxl.load_workbook(review.WORKBOOK)
    check("export: workbook with Read me, Pours, Personalities, Rule candidates, Menu",
          rc == 0 and wb.sheetnames == ["Read me", "Pours", "Personalities", "Rule candidates", "Menu"], out)
    check("export: the personality card sits beside the pour",
          "Sagrada Familia" in (cell("creator-hero", "the personality") or "") and
          "scientist of the senses" in (cell("sage-lover", "the personality") or ""))
    check("export: the Personalities sheet has one full row per pour, both archetypes",
          wb["Personalities"].max_row - 1 == len(before) and
          any("Hero:" in str(c.value) for c in wb["Personalities"][2] if c.value))
    check("export: one row per pour", wb["Pours"].max_row - 1 == len(before), wb["Pours"].max_row)
    check("export: read-only columns locked, words and Robin's columns open",
          wb["Pours"].protection.sheet and wb["Pours"]["A2"].protection.locked
          and not wb["Pours"].cell(row=2, column=[c.value for c in wb["Pours"][1]].index("tagline") + 1).protection.locked)
    check("export: the resonance line is shown", "smaller" in (cell("creator-hero", "this is me") or ""))
    rc, st = run("status")
    check("status: a fresh export shows no changes", rc == 0 and st["changed"] == [], st)

    # 3. sync with no changes touches nothing
    rc, out = run("sync")
    check("sync with no changes: no pour file changes a byte", rc == 0 and hashes() == before and out["pours"] == [], out)

    # 4. Robin edits, approves, sends notes
    edit({"sage-lover": {"tagline": "Some truths last because someone cared enough to keep them."},
          "magician-outlaw": {"yours 2": cell("magician-outlaw", "yours 2") + " One more sentence from Robin.",
                              "Decision": "Approve"},
          "creator-hero": {"Notes to the room": "The ending is too long. Tighten p5."}})
    rc, st = run("status")
    check("status: sees the three touched rows", sorted(st["changed"]) == ["creator-hero", "magician-outlaw", "sage-lover"], st)
    rc, out = run("export")
    check("export: refuses while there are unsynced changes", rc == 1 and "sync" in out.get("reason", ""), out)
    rc, out = run("sync")
    sl = pourfile.parse(os.path.join(POURS, "sage-lover.md"))
    mo = pourfile.parse(os.path.join(POURS, "magician-outlaw.md"))
    check("sync: Robin's tagline is in the pour, verbatim", sl["tagline"] == "Some truths last because someone cared enough to keep them.")
    check("sync: Robin's paragraph edit is in the pour", mo["yours"][1].endswith("One more sentence from Robin."))
    check("sync: Approve sets status approved", mo["status"].startswith("approved"), mo["status"])
    check("sync: a row without Approve stays draft", sl["status"].startswith("draft"), sl["status"])
    room = open(os.path.join(STUDIO, "rooms", "sage-lover.md"), encoding="utf-8").read()
    check("sync: the edit is logged at the top of the room record (created if missing)",
          "## Robin's review" in room and "cared enough" in room and "loved them enough" in room)
    ledger = open(os.path.join(STUDIO, "robin-edits.md"), encoding="utf-8").read()
    check("sync: the edit is in the ledger with its ground", "| sage-lover | tagline |" in ledger and "| Wren |" in ledger)
    index = open(os.path.join(STUDIO, "index.md"), encoding="utf-8").read()
    check("sync: notes queue a rework", out["reworks"] == ["creator-hero"] and "rework queued" in index, out["reworks"])
    ch_room = open(os.path.join(STUDIO, "rooms", "creator-hero.md"), encoding="utf-8").read()
    check("sync: the notes are in the room record for the rework", "Tighten p5." in ch_room.split("## Where things stand")[0])
    rc, st = run("status")
    check("sync: re-exports a clean workbook", rc == 0 and st["changed"] == [], st)
    check("sync: the old workbook is archived", len(os.listdir(os.path.join(STUDIO, "review-archive"))) >= 1)

    # 5. a recipe edit holds the approval (the allergen veto depends on the spec)
    recipe = cell("sage-lover", "recipe")
    edit({"sage-lover": {"recipe": recipe.replace("60 ml | blended scotch", "50 ml | blended scotch"), "Decision": "Approve"}})
    rc, out = run("sync")
    sl = pourfile.parse(os.path.join(POURS, "sage-lover.md"))
    check("recipe edit: written verbatim", sl["recipe"][0]["amount"] == "50 ml")
    check("recipe edit: approval held, status unchanged", sl["status"].startswith("draft") and
          out["pours"][0]["approval"].startswith("held"), out["pours"])
    check("recipe edit: Tomás is asked (desk note)", "Tomás" in (cell("sage-lover", "from the desk") or ""))

    # 6. paragraphs: a blank line splits, an emptied cell removes
    edit({"magician-outlaw": {"yours 1": cell("magician-outlaw", "yours 1") + "\n\nA new paragraph.", "yours 4": ""}})
    n_before = len(pourfile.parse(os.path.join(POURS, "magician-outlaw.md"))["yours"])
    run("sync")
    mo = pourfile.parse(os.path.join(POURS, "magician-outlaw.md"))
    check("paragraphs: split on a blank line and removed when emptied",
          len(mo["yours"]) == n_before and mo["yours"][1] == "A new paragraph.", [p[:30] for p in mo["yours"]])

    # 7. conflict: the pour changed after export -> row skipped, Robin's words kept
    edit({"sage-lover": {"epigraph": "Robin's new epigraph."}})
    p = os.path.join(POURS, "sage-lover.md")
    open(p, "a", encoding="utf-8").write("\n")   # e.g. a rework ran meanwhile
    snapshot = open(p, encoding="utf-8").read()
    rc, out = run("sync")
    check("conflict: the row is skipped and the file untouched", out["conflicts"] == ["sage-lover"] and open(p, encoding="utf-8").read() == snapshot)
    check("conflict: Robin's words survive in the desk notes", "Robin's new epigraph." in (cell("sage-lover", "from the desk") or ""))

    # 8. Excel open -> blocked
    open(review.LOCK, "w").close()
    rc1, o1 = run("sync")
    rc2, o2 = run("export", "--force")
    os.remove(review.LOCK)
    check("Excel open: sync and export both stop", rc1 == 1 and rc2 == 1 and "Excel" in o1["reason"])

    # 9. rule candidates
    cand = os.path.join(STUDIO, "rule-candidates.md")
    open(cand, "a", encoding="utf-8").write("| 2026-09-26 | room (Wren), test | Never end on a question | tested | |\n")
    run("export", "--force")
    edit({"rules": {"Decision": "Accept", "Robin's note": "yes"}})
    rc, out = run("sync")
    check("rules: Accept is recorded and returned for the rulebook",
          "**Accepted**" in open(cand, encoding="utf-8").read().splitlines()[-1]
          and out["rules"]["accepted"][0]["rule"] == "Never end on a question", out["rules"])

    # 10. approve in chat
    set_status("creator-hero", "draft")
    rc, out = run("approve", "creator-hero", "--note", "love it")
    check("approve in chat: status approved and logged",
          pourfile.parse(os.path.join(POURS, "creator-hero.md"))["status"].startswith("approved")
          and "Approve (in chat)" in open(os.path.join(STUDIO, "rooms", "creator-hero.md"), encoding="utf-8").read())
finally:
    shutil.rmtree(TMP, ignore_errors=True)

print("\n%d failure(s)" % fails)
sys.exit(1 if fails else 0)
